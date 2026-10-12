"""Audit the assembled thesis against the format rules in
.agent/skills/rmuti-thesis/SKILL.md, .agent/rules/thesis-format.md and
.agent/rules/chatshop.md. Report only, changes nothing.

Usage:  py check_book.py [book.docx]   (the PDF next to it is read too)
"""
import os
import re
import sys
from collections import defaultdict

import docx
import pymupdf
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")
ROOT = os.path.dirname(os.path.abspath(__file__))
THAI = re.compile(r"[฀-๿]")
# chatshop.md "ห้ามอ้าง" / "คำที่ใช้"
BANNED = ["แชทบอทปัญญาประดิษฐ์", "แชทบอทอัจฉริยะ", "อุปกรณ์เคลื่อนที่", "มือถือ", "สมาร์ทโฟน",
          "vector database", "ฐานข้อมูลเวกเตอร์", "ทีมพัฒนา", "ผู้ร่วมพัฒนา", "พัฒนาโดยผู้อื่น",
          "payment gateway", "WhatsApp", "TikTok", "Shopee", "Lazada", "Horizon", "Vue", "dio ",
          "ประการแรก", "ประการที่สอง", "เกณฑ์ SLA"]


def body_paragraphs(doc):
    for p in doc.paragraphs:
        yield p, False
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    yield p, True


def check_docx(doc, out):
    seen = set()
    in_body = False  # chapter rules apply from บทที่ 1 up to the bibliography
    for p, in_table in body_paragraphs(doc):
        if not in_table and p.style.name == "Heading 1":
            in_body = True
        if not in_table and p.text.strip() == "บรรณานุกรม":
            in_body = False
        if id(p._p) in seen:  # merged cells repeat paragraphs
            continue
        seen.add(id(p._p))
        text = p.text
        if not text.strip():
            continue
        short = text.strip()[:60]
        # 1) Thai text in a run without <w:cs/>: thaiDistribute spreads the letters
        for r in p.runs:
            if THAI.search(r.text) and (r._r.rPr is None or r._r.rPr.find(qn("w:cs")) is None):
                out["Thai run without <w:cs/> (letters stretch)"].append(short)
                break
        # 2) banned words
        for w in BANNED:
            if w in text:
                out[f'banned word "{w}"'].append(short)
        if text.lstrip("\t").startswith("ในปัจจุบัน"):
            out['paragraph starts with "ในปัจจุบัน"'].append(short)
        # 3) double spaces: allowed only after a heading number X.X.X.. or รูปที่/ตารางที่ N.n
        body = re.sub(r"^\s*\d+(\.\d+)+\s\s", "", text)
        body = re.sub(r"(รูปที่|ตารางที่) [\dก-ฮ]+\.\d*\s\s", "", body)  # number may be a SEQ field
        if in_body and "  " in body.strip() and not in_table and "…" not in body and "...." not in body:
            out["double space in text"].append(short)
        # 4) captions need exactly two spaces after the number
        m = re.match(r"(รูปที่|ตารางที่) (\d+\.\d+)( +)", text.strip())
        if m and len(m.group(3)) != 2:
            out["caption: not 2 spaces after number"].append(short)
        # 5) X.X.X headings must not be bold
        if re.match(r"\t*\d+\.\d+\.\d+", text) and any(r.bold for r in p.runs if r.text.strip()):
            out["X.X.X heading in bold"].append(short)
        # 6) hanging indent outside the bibliography
        fi = p.paragraph_format.first_line_indent
        if in_body and fi is not None and fi < 0:
            out["hanging indent (not allowed outside bibliography)"].append(short)


def check_refs(full_text, out):
    """"ดังรูปที่ 3.9" etc. must point to a caption that exists."""
    caps = set(re.findall(r"(รูปที่|ตารางที่) (\d+\.\d+)  ", full_text))
    for kind, num in set(re.findall(r"(?<!\n)(รูปที่|ตารางที่) (\d+\.\d+)(?!  )", full_text)):
        if (kind, num) not in caps:
            out["reference to a missing figure/table"].append(f"{kind} {num}")
    for kind in ("รูปที่", "ตารางที่"):
        nums = sorted({n for k, n in caps if k == kind}, key=lambda s: tuple(map(int, s.split("."))))
        by_ch = defaultdict(list)
        for n in nums:
            c, i = map(int, n.split("."))
            by_ch[c].append(i)
        for c, seq in by_ch.items():
            if seq != list(range(1, len(seq) + 1)):
                out[f"{kind} numbering gap in chapter {c}"].append(str(seq))


def check_pdf(pdf, out):
    d = pymupdf.open(pdf)
    body_start = next(i for i, p in enumerate(d) if p.get_text().strip().startswith("บทที่ 1"))
    full = []
    for i, page in enumerate(d):
        label = f"p{i + 1}"
        t = page.get_text()
        full.append(t)
        lines = []
        for b in page.get_text("dict")["blocks"]:
            for ln in b.get("lines", []):
                s = "".join(sp["text"] for sp in ln["spans"]).strip()
                if s:
                    lines.append((ln["bbox"], s))
        lines = [(bb, s) for bb, s in lines if bb[1] > 95]  # drop page number
        if not lines:
            continue
        lines.sort(key=lambda x: (round(x[0][1]), x[0][0]))
        # 7) heading or table caption left alone at the bottom of a page
        bb, last = lines[-1]
        if i >= body_start and not t.lstrip().startswith("บทที่") and (re.match(r"^\d+(\.\d+)+\s+\S", last) and len(last) < 70
                                or re.match(r"^ตารางที่ \d+\.\d+  ", last)):
            out["heading/caption at page bottom (orphan)"].append(f"{label}: {last[:50]}")
        # 8) one- or two-word last line of a justified paragraph
        if i < body_start:
            continue
        for k in range(1, len(lines) - 1):
            (b0, s0), (b1, s1), (b2, s2) = lines[k - 1], lines[k], lines[k + 1]
            width, full_w = b1[2] - b1[0], b0[2] - b0[0]
            new_para = b2[0] > b1[0] + 20 or re.match(r"^\d+(\.\d+)+\s|^\d+\)|^(รูปที่|ตารางที่)", s2)
            if full_w > 380 and width < 0.25 * 415 and new_para and abs(b1[0] - b0[0]) < 5 \
                    and not re.match(r"^\d+(\.\d+)+\s", s1):
                out["paragraph ends with 1-2 words (<25% line)"].append(f"{label}: …{s0[-25:]} / {s1}")
    return "\n".join(full)


def check_numbering(pdf, out):
    """Page labels top-right and every TOC / list entry against the real page."""
    d = pymupdf.open(pdf)
    labels, firsts = [], []
    for p in d:
        bl = sorted(p.get_text("blocks"), key=lambda b: (b[1], b[0]))
        top = [b[4].strip() for b in bl if b[1] < 95 and b[0] > 400]
        labels.append(top[0] if top else "")
        rest = [b[4].strip() for b in bl if not (b[1] < 95 and b[0] > 400) and b[4].strip()]
        firsts.append(rest[0] if rest else "")
    bs = next(i for i, f in enumerate(firsts) if f.startswith("บทที่ 1"))
    for i, f in enumerate(firsts):
        if not f:
            out["blank page"].append(f"p{i + 1}")
    for i in range(bs, len(d)):
        starts = firsts[i].startswith(("บทที่", "บรรณานุกรม"))
        if starts and labels[i]:
            out["page number shown on a chapter's first page"].append(f"p{i + 1}")
        if not starts and labels[i] != str(i - bs + 1):
            out["wrong arabic page number"].append(f"p{i + 1}: {labels[i]!r}")
    actual, claims = {}, {}
    for i in range(bs, len(d)):
        t = d[i].get_text()
        for m in re.finditer(r"(ตารางที่|รูปที่) (\d+\.\d+)  (.*)", t):
            if not m.group(3).rstrip().endswith("(ต่อ)"):
                actual.setdefault(m.group(1) + m.group(2), i - bs + 1)
        for m in re.finditer(r"^(\d+\.\d+)  ", t, re.M):
            actual.setdefault("H" + m.group(1), i - bs + 1)
        for m in re.finditer(r"^บทที่ (\d)", t, re.M):
            actual.setdefault("C" + m.group(1), i - bs + 1)
        if t.strip().startswith("บรรณานุกรม"):
            actual.setdefault("BIB", i - bs + 1)
    lines = [l.strip() for i in range(bs) for l in d[i].get_text().split("\n") if l.strip()]
    for k, l in enumerate(lines):
        m = re.match(r"(ตารางที่|รูปที่) (\d+\.\d+) ", l)
        key = m and m.group(1) + m.group(2)
        if not key and re.match(r"\d+\.\d+ ", l):
            key = "H" + l.split()[0]
        elif not key and re.match(r"บทที่ \d ", l):
            key = "C" + l.split()[1]
        elif not key and l == "บรรณานุกรม":
            key = "BIB"
        nxt = next((x for x in lines[k + 1:k + 4] if re.fullmatch(r"\d+", x)), None)
        if key and nxt:
            claims.setdefault(key, int(nxt))
    for k, v in claims.items():
        if actual.get(k) != v:
            out["TOC page differs from the real page"].append(f"{k}: TOC {v}, real {actual.get(k)}")
    print(f"numbering: {len(claims)} TOC entries checked, body pages 1..{len(d) - bs}")


def check_figures(doc, out):
    """Every picture has the black outline; the table holding it has no borders."""
    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    PIC = "{http://schemas.openxmlformats.org/drawingml/2006/picture}"
    for k, pic in enumerate(doc.element.body.iter(PIC + "pic"), 1):
        if pic.find(f"{PIC}spPr/{A}ln") is None:
            out["picture without black outline"].append(f"picture #{k}")
    for tbl in doc.element.body.iter(qn("w:tbl")):
        if tbl.find(".//" + PIC + "pic") is None:
            continue
        b = tbl.tblPr.find(qn("w:tblBorders"))
        vals = [] if b is None else [e.get(qn("w:val")) for e in b]
        if b is None or any(v not in ("none", "nil") for v in vals) or tbl.find(".//" + qn("w:tcBorders")) is not None:
            cap = re.search(r"รูปที่ [\d.]+", "".join(tbl.itertext()))
            out["figure box table has borders"].append(cap.group(0) if cap else "?")


def check_citations(full_text, out):
    """Body [n] and bibliography [n] must match both ways."""
    body, _, bib = full_text.rpartition("บรรณานุกรม")
    cited = {int(n) for n in re.findall(r"\[(\d+)\]", body)}
    listed = {int(n) for n in re.findall(r"^\[(\d+)\]", bib, re.M)}
    for n in sorted(cited - listed):
        out["cited [n] missing from the bibliography"].append(str(n))
    for n in sorted(listed - cited):
        out["bibliography entry never cited"].append(str(n))
    print(f"citations: {len(cited)} cited, {len(listed)} in bibliography")


def check_chapter_heads(pdf, out):
    """'บทที่ N' sits at the same height on every chapter's first page."""
    ys = {}
    for i, p in enumerate(pymupdf.open(pdf)):
        for b in p.get_text("dict")["blocks"]:
            for ln in b.get("lines", []):
                s = "".join(sp["text"] for sp in ln["spans"]).strip()
                if re.fullmatch(r"บทที่ \d", s):
                    ys[s] = round(ln["bbox"][1])
    if len(set(ys.values())) > 1:
        out["chapter titles at different heights"].append(str(ys))


def main(argv):
    path = argv[1] if len(argv) > 1 else os.path.join(ROOT, "output", "Chatshop_ปริญญานิพนธ์ฉบับสมบูรณ์_v3.docx")
    out = defaultdict(list)
    doc = docx.Document(path)
    check_docx(doc, out)
    check_figures(doc, out)
    full = check_pdf(path[:-5] + ".pdf", out)
    check_numbering(path[:-5] + ".pdf", out)
    check_chapter_heads(path[:-5] + ".pdf", out)
    check_refs(full, out)
    check_citations(full, out)
    if not out:
        print("OK: no findings")
    for k, v in out.items():
        print(f"\n## {k}  ({len(v)})")
        for x in v[:40]:
            print("  -", x)
        if len(v) > 40:
            print(f"  ... {len(v) - 40} more")


if __name__ == "__main__":
    main(sys.argv)
