import os, re, sys, shutil, copy
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import win32com.client
import pymupdf

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".agent", "skills", "rmuti-thesis", "scripts"))
from thesis_build import split_tables, export_pdf, normalize_figures, _save, _open, unsplit_tables
import ch4_comparison

sys.stdout.reconfigure(encoding='utf-8')

ROOT = os.path.abspath(os.path.dirname(__file__))
TEMP_DIR = os.path.join(ROOT, "process", "temp")
os.makedirs(TEMP_DIR, exist_ok=True)

OUT_DOCX = os.path.join(TEMP_DIR, "Chatshop_บทที่1_5.docx")  # intermediate, kept out of output/
OUT_PDF = os.path.join(TEMP_DIR, "Chatshop_บทที่1_5.pdf")


def src(name):
    """Source chapter files live in output/ or, after a clean-up, output/old/."""
    for d in ("output", os.path.join("output", "old")):
        f = os.path.join(ROOT, d, name)
        if os.path.exists(f):
            return f
    raise FileNotFoundError(name)

def mark_thai(run):
    rPr = run._r.get_or_add_rPr()
    for tag in ("w:rFonts", "w:cs", "w:lang"):
        for old in rPr.findall(qn(tag)):
            rPr.remove(old)
    fonts = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(a), "TH SarabunPSK")
    fonts.set(qn("w:hint"), "cs")
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:eastAsia"), "en-US")
    lang.set(qn("w:bidi"), "th-TH")
    rPr.insert(0, fonts)
    rPr.append(OxmlElement("w:cs"))
    rPr.append(lang)

def prepare_ch3():
    path = src("Chatshop_บทที่3_วิธีดำเนินงาน_v2.docx")
    dst = os.path.join(TEMP_DIR, "ch3_prepared.docx")
    print(f"Preparing Chapter 3 from {path}...")
    doc = docx.Document(path)
    print("  rejoined split tables:", unsplit_tables(doc))  # re-split once on the book layout

    # Remove any internal paragraph-level sectPr (from old senior template page-1 break)
    for p in doc.paragraphs:
        if p._p.pPr is not None:
            for sp in p._p.pPr.findall(qn("w:sectPr")):
                p._p.pPr.remove(sp)

    # Add transition line after 3.8 in the topic list if not present
    for p in doc.paragraphs:
        if p.text.strip() == "3.8 การตั้งค่าแพลตฟอร์มภายนอก":
            nxt = p._p.getnext()
            nxt_txt = docx.text.paragraph.Paragraph(nxt, doc).text.strip() if nxt is not None and nxt.tag == qn("w:p") else ""
            if "ซึ่งแต่ละหัวข้อมีรายละเอียดดังต่อไปนี้" not in nxt_txt:
                ch4_comparison.clone_after(p._p, doc.paragraphs[1], "\tซึ่งแต่ละหัวข้อมีรายละเอียดดังต่อไปนี้")
            break

    # Convert accidental Heading 3 items in Ch3 to Normal with \t3.X.X prefix (RMUTI rule for Ch2+)
    h2_idx = 0
    h3_idx = 0
    for p in doc.paragraphs:
        if p.style.name == "Heading 2":
            h2_idx += 1
            h3_idx = 0
        elif p.style.name == "Normal":
            m = re.match(r"^\t*3\.(\d+)\.(\d+)\s", p.text)
            if m and int(m.group(1)) == h2_idx:
                h3_idx = int(m.group(2))
        elif p.style.name == "Heading 3":
            h3_idx += 1
            p.style = doc.styles["Normal"]
            txt = p.text.lstrip("\t ")
            runs = [r for r in p.runs if r.text]
            if runs:
                runs[0].text = f"\t3.{h2_idx}.{h3_idx} {txt}"
                for r in runs[1:]:
                    r.text = ""
                for r in p.runs:
                    r.bold = None
                    mark_thai(r)

    # Overflowing tables are split after merging (thesis_build.split_tables)
    doc.save(dst)
    print(f"Saved prepared Chapter 3: {dst}")
    return dst

def prepare_ch4():
    path = src("Chatshop_บทที่4_ผลการดำเนินงาน_v2.docx")
    dst = os.path.join(TEMP_DIR, "ch4_prepared.docx")
    print(f"Preparing Chapter 4 from {path}...")
    doc = docx.Document(path)
    print("  rejoined split tables:", unsplit_tables(doc))  # re-split once on the book layout

    # Align Chapter 4 title with friend's passed book: ผลการดำเนินงานและการทดสอบ
    h1 = doc.paragraphs[0]
    if h1.style.name == "Heading 1":
        for r in h1.runs:
            if "ผลการดำเนินงาน" in r.text and "และการทดสอบ" not in r.text:
                r.text = r.text.replace("ผลการดำเนินงาน", "ผลการดำเนินงานและการทดสอบ")

    # Test cases 4.1 to 4.17: Set page_break_before on Heading 2 so each test case starts on a new page cleanly
    for p in doc.paragraphs:
        txt = p.text.strip()
        if p.style.name == "Heading 2" and ("ผลการทดสอบ" in txt or "ข้อผิดพลาด" in txt or "แบบประเมิน" in txt):
            p.paragraph_format.page_break_before = True

    # 4.17 "ข้อผิดพลาดที่พบจากการทดสอบ" (failed cases) is cut: drop the section and its
    # entry in the chapter's topic list; the evaluation section moves up from 4.18 to 4.17
    ps = doc.paragraphs
    start = next(i for i, p in enumerate(ps)
                 if p.style.name == "Heading 2" and p.text.strip() == "ข้อผิดพลาดที่พบจากการทดสอบ")
    end = next(i for i in range(start + 1, len(ps)) if ps[i].style.name == "Heading 2")
    topic = next(p for p in ps[:start] if p.text.strip() == "4.17 ข้อผิดพลาดที่พบจากการทดสอบ")
    for p in [topic] + ps[start:end]:
        assert p._p.pPr is None or p._p.pPr.find(qn("w:sectPr")) is None
        p._p.getparent().remove(p._p)
    renamed = 0
    for p in doc.paragraphs:
        if re.match(r"\s*4\.18( |\.\d)", p.text):
            r = next(r for r in p.runs if "4.18" in r.text)
            r.text = r.text.replace("4.18", "4.17", 1)
            renamed += 1
    assert renamed == 6, f"expected 4.18 + 4.18.1-5 to renumber, got {renamed}"

    # author's fix in v2: table 4.13 Test Steps, step 3 on its own line
    step3 = " 3) แก้ไขคำสั่งซื้อของร้านอื่น"
    runs = [r for t in doc.tables for row in t.rows for c in row.cells
            for cp in c.paragraphs for r in cp.runs if step3 in r.text]
    assert runs, "4.13 Test Steps text changed: update the line-break fix"
    for r in runs:
        r.text = r.text.replace(step3, "\n" + step3.strip())

    ch4_comparison.blank_before_first_table(doc)
    ch4_comparison.add_comparison(doc)  # 4.18 before/after comparison + table 4.23

    # Add transition line after 4.18 in topic list
    for p in doc.paragraphs:
        if p.text.strip().startswith("4.18 การเปรียบเทียบการทำงาน"):
            ch4_comparison.clone_after(p._p, doc.paragraphs[1], "\tซึ่งแต่ละหัวข้อมีรายละเอียดดังต่อไปนี้")
            break

    doc.save(dst)
    print(f"Saved prepared Chapter 4: {dst}")
    return dst

SLIP_SUGGESTION = ("5.4.1 พัฒนาฟังก์ชันตรวจสอบสลิปโอนเงินที่ลูกค้าส่งมาในแชท โดยให้ระบบอ่านยอดเงินและข้อมูล"
                   "การโอนจากรูปสลิป เทียบกับยอดของคำสั่งซื้อที่เปิดไว้ในบทสนทนา แล้วปรับสถานะคำสั่งซื้อเป็นชำระเงินแล้ว "
                   "ช่วยลดเวลาที่เจ้าหน้าที่ต้องตรวจสลิปด้วยตนเอง")


def prepare_ch5():
    path = src("Chatshop_บทที่5_สรุปผลและข้อเสนอแนะ.docx")
    dst = os.path.join(TEMP_DIR, "ch5_prepared.docx")
    print(f"Preparing Chapter 5 from {path}...")
    doc = docx.Document(path)

    # Align Chapter 5 title with friend's passed book: สรุปผลและข้อเสนอแนะ (no space)
    h1 = doc.paragraphs[0]
    if h1.style.name == "Heading 1":
        for r in h1.runs:
            if "สรุปผล และข้อเสนอแนะ" in r.text:
                r.text = r.text.replace("สรุปผล และข้อเสนอแนะ", "สรุปผลและข้อเสนอแนะ")

    # Insert Chapter 5 intro and topic list before 5.1 if not present
    if not any("โดยมีหัวข้อดังนี้" in p.text for p in doc.paragraphs[:6]):
        body_tmpl = next(p for p in doc.paragraphs if p.style.name == "Normal" and p.text.strip())
        anchor = h1._p
        intro_txt = ("\tโครงงานนี้พัฒนาแอปพลิเคชันบนสมาร์ตโฟนและระบบบริการหลังบ้านของ Chatshop ซึ่งรวมแชทจาก "
                     "LINE Official Account, Facebook Messenger, Instagram และ WebChat ไว้ในกล่องข้อความเดียว "
                     "พร้อมระบบแจ้งเตือน Push Notification การควบคุมแชทบอท การจัดการสินค้า ลูกค้า คำสั่งซื้อ และรายงานสรุปผล "
                     "บทนี้สรุปผลการดำเนินงานเทียบกับวัตถุประสงค์ อภิปรายผล ปัญหาและอุปสรรค และข้อเสนอแนะสำหรับการพัฒนาต่อ โดยมีหัวข้อดังนี้")
        p_intro = ch4_comparison.clone_after(anchor, body_tmpl, intro_txt)
        anchor = p_intro._p
        for t_line in ("5.1 สรุปผลการดำเนินงาน", "5.2 อภิปรายผล", "5.3 ปัญหาและอุปสรรค", "5.4 ข้อเสนอแนะและแนวทางในการพัฒนาต่อ"):
            pt = ch4_comparison.clone_after(anchor, body_tmpl, t_line)
            pt.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pt.paragraph_format.left_indent = Cm(1.27)
            anchor = pt._p
        p_trans = ch4_comparison.clone_after(anchor, body_tmpl, "\tซึ่งแต่ละหัวข้อมีรายละเอียดดังต่อไปนี้")
        ch4_comparison.clone_after(p_trans._p, body_tmpl, "")

    # follow the cut of 4.17 (failed test cases) only if the document still has old sections
    has_old_535 = any(p.text.strip().startswith("5.3.5 ") for p in doc.paragraphs)
    if has_old_535:
        def para(prefix):
            return next(p for p in doc.paragraphs if p.text.strip().startswith(prefix))
        for prefix in ("5.1.6 ", "5.3.5 "):
            para(prefix)._p.getparent().remove(para(prefix)._p)
        r = next(r for r in para("5.1.7 ").runs if "5.1.7" in r.text)
        r.text = r.text.replace("5.1.7", "5.1.6", 1)
        p = para("5.4.1 ")
        runs = [r for r in p.runs if r.text]
        runs[-1].text = ("\t" if p.text.startswith("\t") else "") + SLIP_SUGGESTION
        for r in runs[:-1]:
            r.text = ""
        text = "\n".join(p.text for p in doc.paragraphs)
        for gone in ("4.17", "ไม่ผ่าน", "secure_storage", "5.1.7"):
            assert gone not in text, f"chapter 5 still mentions {gone!r}"
        ch4_comparison.update_ch5(doc)  # 5.1.7 comparison summary, 5.4.6 dropped, 5.4.7 measure the gain

    # Verify all runs are tagged Thai
    for p in doc.paragraphs:
        for r in p.runs:
            mark_thai(r)

    doc.save(dst)
    print(f"Saved prepared Chapter 5: {dst}")
    return dst

# bib.txt numbers -> words that, at their first mention in chapters 1 and 3, get "[n]"
# (longer phrases first so "Laravel Reverb" wins over "Laravel")
CITE = [("Laravel Reverb", 12), ("Firebase Cloud Messaging", 13), ("LINE Messaging API", 8),
        ("LINE Official Account", 8), ("Graph API", 9), ("Messenger", 9), ("Flutter", 1), ("Dart", 2),
        ("Riverpod", 3), ("go_router", 4), ("Laravel", 5), ("REST API", 6), ("OAuth", 7), ("Webhook", 10),
        ("คิวงาน", 11), ("MySQL", 14), ("MongoDB", 15), ("หลายผู้เช่า", 17), ("แชทบอท", 19),
        ("แบบจำลองภาษา", 21), ("RAG", 22), ("เวกเตอร์", 23)]


def add_citations(doc, chapters=(1, 3)):
    """First mention of each referenced topic in the given chapters gets its
    IEEE number. Body paragraphs only (no headings, captions or tables)."""
    import re
    chapter, cited, added = 0, set(), []
    for p in doc.paragraphs:
        if p.style.name == "Heading 1":
            chapter, cited = chapter + 1, set()
            continue
        text = p.text.strip()
        if chapter not in chapters or p.style.name.startswith("Heading") or not text \
                or text.startswith(("รูปที่", "ตารางที่")):
            continue
        cited |= {int(n) for n in re.findall(r"\[(\d+)\]", p.text)}
        for word, n in CITE:
            if n in cited:
                continue
            longer = [w for w, _ in CITE if w != word and word in w]
            # not mid-word, not "Laravel 10", not already cited, not inside an English name
            pat = re.escape(word) + r"(?![\w\-])(?!\s*[\[\d])" + (r"(?!\s+[A-Za-z])" if word.isascii() else "")
            for r in p.runs:
                m = next((m for m in re.finditer(pat, r.text)
                          if not any(r.text.startswith(w, m.start() - w.index(word)) for w in longer)
                          and re.sub(r"[\d.\s]", "", r.text[:m.start()])), None)  # skip a term being defined
                if m:
                    r.text = r.text[:m.end()] + f" [{n}]" + r.text[m.end():]
                    cited.add(n)
                    added.append(f"บทที่ {chapter}: {word} [{n}]  …{r.text[max(0, m.start() - 25):m.end() + 6]}…")
                    break
    return added


def autofit_tables(path):
    """Column widths follow the content (as the author set by hand): Word
    AutoFit to contents, then to the page width. Continuation parts "(ต่อ)" copy
    the widths of the part before them. Figure boxes are left alone."""
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        d = word.Documents.Open(os.path.abspath(path))
        widths = None
        for i in range(1, d.Tables.Count + 1):
            t = d.Tables(i)
            cap = t.Range.Paragraphs(1).Range.Previous(4, 1)  # wdParagraph
            cap = cap.Text.strip() if cap is not None else ""
            if not cap.startswith("ตารางที่"):
                continue
            if cap.endswith("(ต่อ)") and widths and len(widths) == t.Columns.Count:
                for k, w in enumerate(widths, 1):
                    t.Columns(k).Width = w
                continue
            t.AllowAutoFit = True
            t.AutoFitBehavior(1)  # wdAutoFitContent
            t.AutoFitBehavior(2)  # wdAutoFitWindow
            widths = [t.Columns(k).Width for k in range(1, t.Columns.Count + 1)]
        d.Save()
        d.Close(False)
    finally:
        word.Quit()

    # Sync table widths with author's v2 if difference exceeds +-0.5 cm
    try:
        v2_ref = src("Chatshop_ปริญญานิพนธ์ฉบับสมบูรณ์_v2.docx")
    except FileNotFoundError:
        v2_ref = None
    if v2_ref:
        v2_doc = _open(v2_ref)
        src_map = {}
        for t in v2_doc.tables:
            p = t._tbl.getprevious()
            txt = docx.text.paragraph.Paragraph(p, v2_doc).text.strip() if p is not None and p.tag.endswith("p") else ""
            if "ตารางที่" in txt:
                grid = t._tbl.xpath(".//w:tblGrid/w:gridCol")
                if grid:
                    src_map[txt] = [int(gc.get(qn("w:w"))) for gc in grid]
        doc = _open(path)
        for t in doc.tables:
            p = t._tbl.getprevious()
            txt = docx.text.paragraph.Paragraph(p, doc).text.strip() if p is not None and p.tag.endswith("p") else ""
            if txt in src_map:
                target_w = src_map[txt]
                grid = t._tbl.xpath(".//w:tblGrid/w:gridCol")
                # v2 tables the author never resized still carry the generator's widths
                # (4 cols 3.6/2.9/4.7/3.45 cm, 2 cols 50/50): keep autofit for those
                untouched = (len(target_w) == 2 and abs(target_w[0] - target_w[1]) < 60) or \
                    (len(target_w) == 4 and all(abs(a - b) < 60 for a, b in zip(target_w, (2041, 1644, 2660, 1951))))
                if grid and len(grid) == len(target_w) and not untouched:
                    cur_w = [int(gc.get(qn("w:w"))) for gc in grid]
                    if any(abs(a - b) > 283 for a, b in zip(cur_w, target_w)):  # 0.5 cm = 283 dxa
                        for gc, tw in zip(grid, target_w):
                            gc.set(qn("w:w"), str(tw))
                        for row in t.rows:
                            for cell, tw in zip(row.cells, target_w):
                                cell.width = docx.shared.Length(tw * 635)
        _save(doc, path)

    # satisfaction tables 4.18-4.22 all get the widths the author set on 4.19
    doc = _open(path)
    for t in doc.tables:
        if len(t.columns) == 4 and t.rows[0].cells[0].text.strip() in ("รายการประเมิน", "ด้านการประเมิน"):
            set_widths(t, EVAL_CM)
    _save(doc, path)


EVAL_CM = (7.29, 2.06, 2.86, 2.42)  # author's widths on table 4.19 (v2)


def set_widths(t, cms):
    t.autofit = False  # tblLayout fixed: Word must not re-autofit these columns
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), cms):
        gc.set(qn("w:w"), str(round(w * 567)))
    for row in t.rows:
        for cell, w in zip(row.cells, cms):
            cell.width = Cm(w)


def main():
    print("=" * 60)
    print("Step 1: Preparing chapters with strict thesis rules...")
    print("=" * 60)
    
    # Ensure Chapter 1-2 is up to date
    ch1_2_docx = src("Chatshop_บทที่1_2.docx")
        
    ch3_docx = prepare_ch3()
    ch4_docx = prepare_ch4()
    ch5_docx = prepare_ch5()

    print("\n" + "=" * 60)
    print("Step 2: Merging Chapters 1-5 via Word COM...")
    print("=" * 60)

    print(f"Base file: {ch1_2_docx} -> {OUT_DOCX}")
    shutil.copy2(ch1_2_docx, OUT_DOCX)

    word = win32com.client.DispatchEx("Word.Application")  # own instance, never the user's open Word
    word.Visible = False
    word.DisplayAlerts = 0  # no hidden modal dialogs
    try:
        doc = word.Documents.Open(os.path.abspath(OUT_DOCX))

        # Merge Chapter 3
        print("Appending Chapter 3...")
        rng = doc.Content
        rng.Collapse(0)  # wdCollapseEnd
        rng.InsertBreak(2)  # wdSectionBreakNextPage
        rng.Collapse(0)
        rng.InsertFile(os.path.abspath(ch3_docx))

        # Merge Chapter 4
        print("Appending Chapter 4...")
        rng = doc.Content
        rng.Collapse(0)  # wdCollapseEnd
        rng.InsertBreak(2)  # wdSectionBreakNextPage
        rng.Collapse(0)
        rng.InsertFile(os.path.abspath(ch4_docx))

        # Merge Chapter 5
        print("Appending Chapter 5...")
        rng = doc.Content
        rng.Collapse(0)  # wdCollapseEnd
        rng.InsertBreak(2)  # wdSectionBreakNextPage
        rng.Collapse(0)
        rng.InsertFile(os.path.abspath(ch5_docx))

        print("Enforcing section margins on all sections...")
        for i in range(1, doc.Sections.Count + 1):
            sec = doc.Sections(i)
            # RMUTI thesis margin: Top 3.81 cm (108 pt), Bottom 2.54 cm (72 pt), Left 3.81 cm (108 pt), Right 2.54 cm (72 pt)
            sec.PageSetup.TopMargin = 108
            sec.PageSetup.BottomMargin = 72
            sec.PageSetup.LeftMargin = 108
            sec.PageSetup.RightMargin = 72

        print("Updating fields and repaginating...")
        doc.Fields.Update()
        doc.Repaginate()
        doc.Save()
        doc.Close(False)
    finally:
        word.Quit()

    # Word can open a file in recovery mode and then silently not save it: check before going on
    merged = _open(OUT_DOCX)
    chapters = sum(p.style.name == "Heading 1" for p in merged.paragraphs)
    assert chapters == 5, f"merged file has {chapters} chapters, expected 5 (Word did not save the merge)"
    for p in merged.paragraphs:  # a heading never sits alone at the bottom of a page
        if p.style.name in ("Heading 1", "Heading 2"):
            p.paragraph_format.keep_with_next = True
        if p.style.name == "Heading 1":  # every chapter title at the same height (ch3 lacked it)
            p.paragraph_format.space_before = Pt(36)
    normalize_figures(merged)
    for line in add_citations(merged):
        print("cite:", line)
    _save(merged, OUT_DOCX)  # Word may still hold the file for a moment
    print("Fitting table columns to their content...")
    autofit_tables(OUT_DOCX)
    print("Splitting tables that run past a page end...")
    split_tables(OUT_DOCX)
    print(f"Exporting PDF: {OUT_PDF}...")
    export_pdf(os.path.abspath(OUT_DOCX), os.path.abspath(OUT_PDF))

    print("\n" + "=" * 60)
    print("Step 3: Rendering key verification pages...")
    print("=" * 60)
    pdf = pymupdf.open(OUT_PDF)
    print(f"Total PDF pages: {len(pdf)}")

    # Search for specific pages
    p_fig2_1 = None
    p_tbl3_1_cont = None
    p_tbl4_1 = None
    p_tbl4_1_cont = None
    p_ch5_start = None

    for i, page in enumerate(pdf):
        txt = page.get_text()
        if not p_fig2_1 and "รูปที่ 2.1 " in txt:
            p_fig2_1 = i + 1
        if not p_tbl3_1_cont and "ตารางที่ 3.1" in txt and "(ต่อ)" in txt:
            p_tbl3_1_cont = i + 1
        if not p_tbl4_1 and "ตารางที่ 4.1" in txt and "(ต่อ)" not in txt:
            p_tbl4_1 = i + 1
        if not p_tbl4_1_cont and "ตารางที่ 4.1" in txt and "(ต่อ)" in txt:
            p_tbl4_1_cont = i + 1
        if not p_ch5_start and "บทที่ 5" in txt and "สรุปผล และข้อเสนอแนะ" in txt:
            p_ch5_start = i + 1

    print(f"Key pages found: Fig 2.1=p{p_fig2_1}, Tbl 3.1(ต่อ)=p{p_tbl3_1_cont}, Tbl 4.1=p{p_tbl4_1}, Tbl 4.1(ต่อ)=p{p_tbl4_1_cont}, Ch5 Start=p{p_ch5_start}")
    
    render_pages = [p for p in [1, p_fig2_1, p_tbl3_1_cont, p_tbl4_1, p_tbl4_1_cont, p_ch5_start, len(pdf)] if p]
    for pno in render_pages:
        if 1 <= pno <= len(pdf):
            img_path = os.path.join(TEMP_DIR, f"master_ch1_5_p{pno}.png")
            pdf[pno - 1].get_pixmap(dpi=130).save(img_path)
            print(f"Rendered page {pno} -> {img_path}")

    print("\nMaster document build complete!")

if __name__ == "__main__":
    main()
