"""
build_v7_book.py
เวอร์ชันแก้ไขตามคำสั่งผู้ใช้:
1. นำ Chatshop_ปริญญานิพนธ์ฉบับสมบูรณ์_v6.docx ที่ผู้ใช้แก้แล้วเป็นต้นฉบับ
2. กิตติกรรมประกาศ: คงเดิมตามที่ผู้ใช้แก้ไว้ 100% ห้ามแตะต้อง ห้ามแปลงเป็นตาราง
3. คำสำคัญ: คงเดิมตามที่ผู้ใช้แก้ไว้ 100%
4. สารบัญตารางและสารบัญรูป:
   - ปรับ Font size จาก 20pt ให้เป็น TH SarabunPSK 16pt ทั้งหมด (แก้ไข style table of figures ลบ szCs=40 และ Angsana New)
   - รายการทั่วไปที่พอดี 1 บรรทัด ให้คงเดิมเป็น 1 บรรทัด (ไม่ไปตัดมั่ว)
   - เฉพาะรายการที่ยาวจนล้นจริง ๆ เท่านั้น (โดยเฉพาะ ตารางที่ 4.4):
     ย้ายคำที่เกินลงมาบรรทัดที่สอง เพื่อให้บรรทัดแรกมีที่ว่างพอสำหรับเลขหน้าชิดขวา
     เช่น ตารางที่ 4.4:
       Line 1: ตารางที่ 4.4  การทดสอบการตรวจลายเซ็นของ Webhook จาก LINE   [Tab]  72
       Line 2: และ Meta
5. ส่งออกเป็น v7.docx และ v7.pdf ผ่าน Word COM
"""

import os
import sys
import shutil
import docx

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

from docx.shared import Pt
from docx.oxml import parse_xml
from docx.oxml.ns import qn, nsdecls

V6_PATH = r"D:\Thesis\doc\output\Chatshop_ปริญญานิพนธ์ฉบับสมบูรณ์_v6.docx"
V7_DOCX = r"D:\Thesis\doc\output\Chatshop_ปริญญานิพนธ์ฉบับสมบูรณ์_v7.docx"
V7_PDF  = r"D:\Thesis\doc\output\Chatshop_ปริญญานิพนธ์ฉบับสมบูรณ์_v7.pdf"

def set_font_16(r, bold=False):
    """กำหนด Run ให้เป็น TH SarabunPSK 16pt แน่นอน 100%"""
    r.font.name = "TH SarabunPSK"
    r.font.size = Pt(16)
    r.bold = bold
    rPr = r._r.get_or_add_rPr()
    
    # rFonts
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="TH SarabunPSK" w:hAnsi="TH SarabunPSK" w:cs="TH SarabunPSK"/>')
        rPr.append(rFonts)
    else:
        rFonts.set(qn('w:ascii'), "TH SarabunPSK")
        rFonts.set(qn('w:hAnsi'), "TH SarabunPSK")
        rFonts.set(qn('w:cs'), "TH SarabunPSK")
        
    # sz & szCs = 32 (16pt)
    sz = rPr.find(qn('w:sz'))
    if sz is None:
        sz = parse_xml(f'<w:sz {nsdecls("w")} w:val="32"/>')
        rPr.append(sz)
    else:
        sz.set(qn('w:val'), "32")
        
    szCs = rPr.find(qn('w:szCs'))
    if szCs is None:
        szCs = parse_xml(f'<w:szCs {nsdecls("w")} w:val="32"/>')
        rPr.append(szCs)
    else:
        szCs.set(qn('w:val'), "32")

def fix_styles(doc):
    """แก้ไข style table of figures และ toc ให้เป็น TH SarabunPSK 16pt ทั้งหมด ลบ Angsana New และ 20pt ทิ้ง"""
    for s in doc.styles:
        s_name = s.name.lower()
        if 'table of figures' in s_name or 'toc' in s_name or s.style_id == 'af9':
            rPr = s.element.find(qn('w:rPr'))
            if rPr is not None:
                # แก้ rFonts
                rFonts = rPr.find(qn('w:rFonts'))
                if rFonts is not None:
                    rFonts.set(qn('w:ascii'), "TH SarabunPSK")
                    rFonts.set(qn('w:hAnsi'), "TH SarabunPSK")
                    rFonts.set(qn('w:cs'), "TH SarabunPSK")
                else:
                    rPr.append(parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="TH SarabunPSK" w:hAnsi="TH SarabunPSK" w:cs="TH SarabunPSK"/>'))
                
                # แก้ sz & szCs
                sz = rPr.find(qn('w:sz'))
                if sz is not None:
                    sz.set(qn('w:val'), "32")
                else:
                    rPr.append(parse_xml(f'<w:sz {nsdecls("w")} w:val="32"/>'))
                    
                szCs = rPr.find(qn('w:szCs'))
                if szCs is not None:
                    szCs.set(qn('w:val'), "32")
                else:
                    rPr.append(parse_xml(f'<w:szCs {nsdecls("w")} w:val="32"/>'))

# เฉพาะรายการที่ยาวจนล้นบรรทัดจริง ๆ เท่านั้น
# Key: รหัสข้อความเป้าหมาย, Value: (Line 1 text, Line 2 text)
SPLIT_RULES_EXACT = {
    # ตารางที่ 4.4 ตามที่ผู้ใช้สั่งเป๊ะ: เอาคำว่า "และ" ย้ายลงไปขึ้นบรรทัดใหม่เป็น "และ Meta" เพื่อให้มีที่ว่างใส่เลข 72 ชิดขวาที่บรรทัดแรก
    "ตารางที่ 4.4": ("ตารางที่ 4.4  การทดสอบการตรวจลายเซ็นของ Webhook จาก LINE", "และ Meta"),
    
    # ตารางที่ชื่อยาวมากในหน้า 11 ที่ล้นจริง
    "ตารางที่ 4.18": ("ตารางที่ 4.18  ผลการประเมินความพึงพอใจด้านฟังก์ชันการทำงาน", "และความสามารถของระบบ"),
    "ตารางที่ 4.19": ("ตารางที่ 4.19  ผลการประเมินความพึงพอใจด้านการออกแบบส่วนต่อประสาน", "และความง่ายในการใช้งาน"),
    "ตารางที่ 4.20": ("ตารางที่ 4.20  ผลการประเมินความพึงพอใจด้านประสิทธิภาพ", "และความเสถียรของระบบ"),
    "ตารางที่ 4.21": ("ตารางที่ 4.21  ผลการประเมินความพึงพอใจด้านประโยชน์", "และความพึงพอใจในภาพรวม"),
    "ตารางที่ 4.23": ("ตารางที่ 4.23  การเปรียบเทียบการทำงานก่อนและหลังใช้", "แอปพลิเคชัน Chatshop"),

    # รูปที่ชื่อยาวมากที่ล้นจริง
    "รูปที่ 2.10": ("รูปที่ 2.10  การทำงานของ Webhook", "เทียบกับการสอบถามข้อมูลเป็นระยะ"),
    "รูปที่ 2.13": ("รูปที่ 2.13  การส่ง Push Notification ผ่าน", "Firebase Cloud Messaging"),
    "รูปที่ 3.26": ("รูปที่ 3.26  ขั้นตอนการรับเหตุการณ์และการแจ้งเตือน", "บนแอปพลิเคชัน"),
    "รูปที่ 3.49": ("รูปที่ 3.49  Channel ID และ Channel Secret", "ใน LINE Developers Console"),
    "รูปที่ 3.51": ("รูปที่ 3.51  การกรอก Channel ID และ Channel Secret", "ในแอปพลิเคชัน"),
    "รูปที่ 3.54": ("รูปที่ 3.54  การตั้งค่า Valid OAuth Redirect URIs", "ของแอปเชื่อมต่อช่องทาง"),
    "รูปที่ ก.19": ("รูปที่ ก.19  การกรอก Channel ID และ Channel Secret", "ในแอปพลิเคชัน"),
}

def fix_toc_lot_lof(doc):
    """
    ปรับแต่งสารบัญ สารบัญตาราง สารบัญรูป:
    - รายการทั่วไป: คงเดิมเป็น 1 บรรทัด ฟอนต์ TH SarabunPSK 16pt เลขหน้าชิดขวา
    - รายการใน SPLIT_RULES_EXACT: ย้ายส่วนที่ล้นลงบรรทัดที่สอง เลขหน้าชิดขวาที่บรรทัดแรก
    - ปรับทุก Run ให้เป็น TH SarabunPSK 16pt (ไม่ให้เป็น 20pt อีกต่อไป)
    """
    start_idx = 86
    end_idx = 290
    
    for idx, p in enumerate(doc.paragraphs[:350]):
        txt = p.text.strip()
        if txt == "สารบัญ":
            start_idx = idx
        if "บทนำ" in txt and len(txt) < 15 and idx > 250:
            end_idx = idx
            break
            
    print(f"Fixing TOC/LOT/LOF from paragraph {start_idx} to {end_idx}...")
    
    for i in range(start_idx, end_idx):
        p = doc.paragraphs[i]
        txt = p.text.strip()
        if not txt:
            continue
            
        # ข้ามหัวข้อหลักของหน้า
        is_heading = any(txt == h for h in [
            "สารบัญ", "สารบัญ (ต่อ)", "สารบัญตาราง", "สารบัญตาราง (ต่อ)", "สารบัญรูป", "สารบัญรูป (ต่อ)",
            "หน้า", "ตาราง\tหน้า", "รูป\tหน้า", "บทที่\tหน้า"
        ])
        if is_heading:
            for r in p.runs:
                set_font_16(r, bold=True)
            continue

        parts = txt.split('\t')
        if len(parts) < 2:
            for r in p.runs:
                set_font_16(r, bold=False)
            continue
            
        full_title = parts[0].strip()
        page_num = parts[-1].strip()
        
        # ตรวจสอบว่าตรงกับรายการที่ต้องตัดบรรทัดหรือไม่
        matched_key = None
        for key in sorted(SPLIT_RULES_EXACT.keys(), key=len, reverse=True):
            if full_title.startswith(key):
                matched_key = key
                break
                
        # ตรวจสอบ Tab right ใน pPr
        pPr = p._p.get_or_add_pPr()
        tabs = pPr.find(qn('w:tabs'))
        if tabs is None:
            tabs = parse_xml(f'<w:tabs {nsdecls("w")}><w:tab w:val="right" w:pos="8296"/></w:tabs>')
            pPr.append(tabs)
        else:
            has_right = any(tab.get(qn('w:val')) == 'right' for tab in tabs.findall(qn('w:tab')))
            if not has_right:
                tabs.append(parse_xml(f'<w:tab {nsdecls("w")} w:val="right" w:pos="8296"/>'))
                
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0

        if matched_key:
            line1_text, line2_text = SPLIT_RULES_EXACT[matched_key]
            print(f"  P[{i}] Splitting multiline: '{line1_text}' | page={page_num} | '{line2_text}'")
            p.text = ""
            
            # Line 1: ข้อความส่วนแรก + Tab + เลขหน้า (ชิดขวาที่ Line 1)
            r1 = p.add_run(line1_text)
            set_font_16(r1, bold=False)
            r2 = p.add_run(f"\t{page_num}")
            set_font_16(r2, bold=False)
            
            # Line 2: ข้อความส่วนที่เหลือ (ไม่มีเลขหน้า)
            r3 = p.add_run()
            r3._r.append(parse_xml(f'<w:br {nsdecls("w")}/>'))
            set_font_16(r3, bold=False)
            r4 = p.add_run(line2_text)
            set_font_16(r4, bold=False)
        else:
            # รายการบรรทัดเดียวปกติ คงเดิมทุกอย่าง แค่ตั้งฟอนต์ 16pt ให้ถูกต้อง
            p.text = ""
            r1 = p.add_run(full_title)
            set_font_16(r1, bold=False)
            r2 = p.add_run(f"\t{page_num}")
            set_font_16(r2, bold=False)

def export_docx_to_pdf(docx_path, pdf_path):
    """ส่งออกเป็น PDF ผ่าน Microsoft Word COM API อย่างปลอดภัย"""
    import win32com.client
    import pythoncom

    pythoncom.CoInitialize()
    word = None
    try:
        print(f"Starting Microsoft Word COM to export: {pdf_path}...")
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = False
        
        doc = word.Documents.Open(os.path.abspath(docx_path))
        doc.ExportAsFixedFormat(
            OutputFileName=os.path.abspath(pdf_path),
            ExportFormat=17, # wdExportFormatPDF
            OpenAfterExport=False,
            OptimizeFor=0, # wdExportOptimizeForPrint
            CreateBookmarks=1 # wdExportCreateHeadingBookmarks
        )
        doc.Close(SaveChanges=False)
        print("PDF Exported successfully!")
    finally:
        if word is not None:
            try:
                word.Quit(SaveChanges=False)
            except Exception:
                pass
        pythoncom.CoUninitialize()

def main():
    print("=== Step 1: Duplicating v6.docx to v7.docx ===")
    shutil.copyfile(V6_PATH, V7_DOCX)
    print(f"Copied {V6_PATH} to {V7_DOCX}")
    
    print("\n=== Step 2: Modifying v7.docx ===")
    doc = docx.Document(V7_DOCX)
    
    print("Fixing styles (table of figures & toc -> TH SarabunPSK 16pt)...")
    fix_styles(doc)
    
    # หมายเหตุ: กิตติกรรมประกาศ และ คำสำคัญ ไม่แตะต้องเด็ดขาด คงเดิมตามที่ user แก้ไว้ใน v6
    
    print("Fixing TOC, LOT, and LOF entries (only split truly overflowing entries)...")
    fix_toc_lot_lof(doc)
    
    print(f"\nSaving updated document to {V7_DOCX}...")
    doc.save(V7_DOCX)
    print("Saved v7.docx successfully!")
    
    print("\n=== Step 3: Exporting v7.pdf via Word COM ===")
    export_docx_to_pdf(V7_DOCX, V7_PDF)
    print("All steps completed successfully!")

if __name__ == "__main__":
    main()
