import os, sys, shutil
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Pt, Cm
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
import win32com.client
import pymupdf

sys.stdout.reconfigure(encoding='utf-8')

ROOT = os.path.abspath(os.path.dirname(__file__))
OLD_V5 = os.path.join(ROOT, "output", "old", "Chatshop_ปริญญานิพนธ์ฉบับสมบูรณ์_v5.docx")
NEW_V5_DOCX = os.path.join(ROOT, "output", "Chatshop_ปริญญานิพนธ์ฉบับสมบูรณ์_v5.docx")
NEW_V5_PDF = os.path.join(ROOT, "output", "Chatshop_ปริญญานิพนธ์ฉบับสมบูรณ์_v5.pdf")

def mark_thai(run, font_size=16, bold=False, italic=False):
    run.bold = bold
    run.italic = italic
    rPr = run._r.get_or_add_rPr()
    for tag in ("w:rFonts", "w:cs", "w:sz", "w:szCs", "w:lang"):
        for old in rPr.findall(qn(tag)):
            rPr.remove(old)
    fonts = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(a), "TH SarabunPSK")
    fonts.set(qn("w:hint"), "cs")
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:eastAsia"), "en-US")
    lang.set(qn("w:bidi"), "th-TH")
    
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(font_size * 2)))
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), str(int(font_size * 2)))
    
    rPr.insert(0, fonts)
    rPr.append(sz)
    rPr.append(szCs)
    rPr.append(OxmlElement("w:cs"))
    rPr.append(lang)

def set_p_text(p, text, align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, size=16, space_b=0, space_a=0, line_spacing=1.0):
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(space_b)
    pf.space_after = Pt(space_a)
    pf.line_spacing = line_spacing
    
    for r in list(p.runs):
        p._p.remove(r._r)
        
    if text:
        r = p.add_run(text)
        mark_thai(r, font_size=size, bold=bold)
    return p

def set_header_p(p, text):
    """Formats abstract metadata headers with exact 4.0 cm hanging indent & tab stop,
    matching reference '01 ปกนอก ถึง กิตติกรรมประกาศ.docx' exactly.
    """
    for r in list(p.runs):
        p._p.remove(r._r)
    if p._p.pPr is not None:
        p._p.remove(p._p.pPr)
    
    pPr_xml = (
        f'<w:pPr {nsdecls("w")}>'
        '<w:pStyle w:val="FrontNormal"/>'
        '<w:tabs><w:tab w:val="left" w:pos="2268"/></w:tabs>'
        '<w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>'
        '<w:ind w:left="2268" w:hanging="2268"/>'
        '<w:jc w:val="left"/>'
        '</w:pPr>'
    )
    p._p.append(parse_xml(pPr_xml))
    r = p.add_run(text)
    mark_thai(r, font_size=16)
    return p

def update_book():
    print(f"Loading base book from: {OLD_V5}")
    shutil.copy2(OLD_V5, NEW_V5_DOCX)
    doc = docx.Document(NEW_V5_DOCX)
    
    # Fix FrontNormal style to not add 8pt space_after and 1.08 line_spacing
    if "FrontNormal" in doc.styles:
        st = doc.styles["FrontNormal"]
        st.paragraph_format.space_after = Pt(0)
        st.paragraph_format.space_before = Pt(0)
        st.paragraph_format.line_spacing = 1.0
        print("Updated FrontNormal style spacing to 0pt / single line spacing.")

    # Group paragraphs by section
    sec_paras = {i: [] for i in range(len(doc.sections))}
    sec_idx = 0
    for p in doc.paragraphs:
        sec_paras[sec_idx].append(p)
        if p._p.pPr is not None and p._p.pPr.find(qn('w:sectPr')) is not None:
            sec_idx += 1

    # 1. Approval page (Section 3) - ensure 2 spaces in author name
    s3 = sec_paras[3]
    for p in s3:
        if "จัดทำโดย" in p.text and "ธนพล วุฒิธาดา" in p.text:
            p.text = p.text.replace("ธนพล วุฒิธาดา", "ธนพล  วุฒิธาดา")
            for r in p.runs: mark_thai(r)

    # 2. Section 4: Thai Abstract (หน้า ง)
    s4 = sec_paras[4]
    print(f"Updating Section 4 (Thai Abstract, {len(s4)} paragraphs)...")
    
    set_header_p(s4[0], "หัวข้อปริญญานิพนธ์\tแอปพลิเคชันรวมศูนย์ข้อความและการจัดการคำสั่งซื้อสำหรับธุรกิจออนไลน์ (Chatshop)")
    set_header_p(s4[1], "จัดทำโดย\tธนพล  วุฒิธาดา")
    set_header_p(s4[2], "ปีที่ปริญญานิพนธ์สำเร็จ\tพ.ศ. 2569")
    set_header_p(s4[3], "สาขาวิชา\tวิศวกรรมคอมพิวเตอร์")
    set_header_p(s4[4], "ประธานที่ปรึกษา\tผู้ช่วยศาสตราจารย์ ดร.อติราช  สุขสวัสดิ์")

    # Empty spacer lines before "บทคัดย่อ": exactly 2 blank lines (16pt) as per RMUTI manual
    set_p_text(s4[5], "") # blank line 1
    set_p_text(s4[6], "") # blank line 2
    # Clear redundant blank lines (s4[7], s4[8])
    set_p_text(s4[7], "")
    s4[7].paragraph_format.line_spacing = Pt(0.01) # collapse
    set_p_text(s4[8], "")
    s4[8].paragraph_format.line_spacing = Pt(0.01) # collapse

    # "บทคัดย่อ" heading: centered bold 16pt
    set_p_text(s4[9], "บทคัดย่อ", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)

    # 1 blank line after "บทคัดย่อ"
    set_p_text(s4[10], "")

    # Abstract body paragraphs
    abs_th_1 = (
        "\tโครงงานนี้นำเสนอการพัฒนาแอปพลิเคชันรวมศูนย์ข้อความและการจัดการคำสั่งซื้อสำหรับธุรกิจออนไลน์ (Chatshop) "
        "เพื่อแก้ไขปัญหาความยุ่งยากในการบริหารจัดการข้อความของร้านค้าออนไลน์ที่ติดต่อสื่อสารกับลูกค้าผ่านหลากหลายแพลตฟอร์ม "
        "อาทิ LINE Official Account, Facebook Messenger, Instagram และ WebChat ซึ่งทำให้เกิดปัญหาการตอบกลับล่าช้า "
        "ข้อความตกหล่น และข้อมูลคำสั่งซื้อกระจัดกระจาย โดยระบบประกอบด้วยระบบบริการหลังบ้านที่พัฒนาด้วย Laravel 10 "
        "ร่วมกับฐานข้อมูล MySQL และ MongoDB สำหรับประมวลผล Webhook ขาเข้า และแอปพลิเคชันบนสมาร์ตโฟนที่พัฒนาด้วย Flutter "
        "ที่ช่วยให้เจ้าหน้าที่สามารถติดตามและตอบกลับข้อความจากทุกช่องทางได้ภายในหน้าต่างเดียวแบบเรียลไทม์"
    )
    set_p_text(s4[11], abs_th_1, align=WD_ALIGN_PARAGRAPH.THAI_JUSTIFY)

    abs_th_2 = (
        "\tนอกจากนี้ระบบยังได้บูรณาการระบบแชทบอทที่รองรับการค้นคืนข้อมูลจากฐานความรู้ (Retrieval-Augmented Generation: RAG) "
        "เพื่อช่วยแนะนำคำตอบและตอบกลับลูกค้าอัตโนมัติ พร้อมทั้งฟังก์ชันการแจ้งเตือนแชทที่ค้างตอบเกินกำหนด "
        "และการเปิดคำสั่งซื้อสินค้าได้โดยตรงจากหน้าต่างสนทนา จากผลการทดสอบระบบและการประเมินความพึงพอใจจากกลุ่มผู้ใช้งานจริงจำนวน 15 คน "
        "พบว่าระบบได้รับความพึงพอใจในภาพรวมอยู่ในระดับดี โดยมีค่าเฉลี่ยรวมเท่ากับ 4.10 และค่าส่วนเบี่ยงเบนมาตรฐานเท่ากับ 0.37 "
        "ซึ่งแสดงให้เห็นว่าระบบสามารถช่วยเพิ่มประสิทธิภาพและลดภาระในการบริหารจัดการร้านค้าออนไลน์ได้อย่างเป็นรูปธรรม"
    )
    set_p_text(s4[12], abs_th_2, align=WD_ALIGN_PARAGRAPH.THAI_JUSTIFY)

    set_p_text(s4[13], "คำสำคัญ\t\t:  แอปพลิเคชันสมาร์ตโฟน, รวมศูนย์ข้อความ, จัดการคำสั่งซื้อ, แชทบอท, Flutter, Laravel")
    s4[13].paragraph_format.tab_stops.add_tab_stop(Cm(1.8), WD_TAB_ALIGNMENT.LEFT)
    
    # Clear any extra trailing blank lines in s4
    for idx in range(14, len(s4)):
        set_p_text(s4[idx], "")
        s4[idx].paragraph_format.line_spacing = Pt(0.01)

    # 3. Section 5: English Abstract (หน้า จ)
    s5 = sec_paras[5]
    print(f"Updating Section 5 (English Abstract, {len(s5)} paragraphs)...")
    
    set_header_p(s5[0], "Project Title\tCentralized Chat and Order Management")
    set_header_p(s5[1], "\tApplication for Online Business")
    set_header_p(s5[2], "Proposed by\tThanaphol  Wuttitada")
    set_header_p(s5[3], "Year\t2026")
    set_header_p(s5[4], "Department\tComputer Engineering")
    set_header_p(s5[5], "Advisor\tAssistant Professor, Atirarj Suksawad, D.Eng")

    # Exactly 2 blank lines before "Abstract"
    set_p_text(s5[6], "") # blank line 1
    set_p_text(s5[7], "") # blank line 2
    # Clear redundant blank line s5[8]
    set_p_text(s5[8], "")
    s5[8].paragraph_format.line_spacing = Pt(0.01)

    # "Abstract" heading: centered bold 16pt
    set_p_text(s5[9], "Abstract", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)

    # 1 blank line after "Abstract"
    set_p_text(s5[10], "")

    abs_en_1 = (
        "\tThis project presents the development of Chatshop, a centralized messaging and order management mobile application for online businesses. "
        "The system aims to address the operational challenges of managing omnichannel customer inquiries across platforms such as LINE Official Account, "
        "Facebook Messenger, Instagram, and WebChat, which typically lead to delayed responses and fragmented order data. "
        "The architecture comprises a robust backend service developed using Laravel 10 with MySQL and MongoDB databases for real-time webhook ingestion, "
        "and a cross-platform mobile application developed using Flutter, enabling store staff to manage all communications within a single unified inbox."
    )
    set_p_text(s5[11], abs_en_1, align=WD_ALIGN_PARAGRAPH.THAI_JUSTIFY)

    abs_en_2 = (
        "\tFurthermore, the system integrates a chatbot equipped with Retrieval-Augmented Generation (RAG) to provide automated responses "
        "and suggested replies, along with overdue response alert notifications and integrated chat-order creation features. "
        "The functional test results and user satisfaction evaluation conducted with 15 representative users indicated a high level of satisfaction, "
        "with an overall mean score of 4.10 and a standard deviation of 0.37. "
        "These results demonstrate that Chatshop significantly enhances workflow efficiency and streamlines communication for modern online merchants."
    )
    set_p_text(s5[12], abs_en_2, align=WD_ALIGN_PARAGRAPH.THAI_JUSTIFY)

    set_p_text(s5[13], "Keywords\t\t:  Mobile Application, Centralized Messaging, Order Management, Chatbot, Flutter, Laravel")
    s5[13].paragraph_format.tab_stops.add_tab_stop(Cm(1.8), WD_TAB_ALIGNMENT.LEFT)

    for idx in range(14, len(s5)):
        set_p_text(s5[idx], "")
        s5[idx].paragraph_format.line_spacing = Pt(0.01)

    # 4. Chapter 4: Page break before Table 4.1 so 4.1-4.18 remain on the first page leaving the bottom blank
    ch4_table_break_set = False
    for p in doc.paragraphs:
        if 'ตารางที่ 4.' in p.text and 'ขอบเขตข้อมูลที่ใช้ในตารางกรณีทดสอบ' in p.text and '(ต่อ)' not in p.text:
            p.paragraph_format.page_break_before = True
            ch4_table_break_set = True
            print("Set page_break_before = True on Chapter 4 Table 4.1.")
            break
    if not ch4_table_break_set:
        print("WARNING: Table 4.1 heading paragraph not found!")

    # 5. Author Biography (ประวัติผู้เขียน):
    # a. Remove '[รูปถ่าย 1.5–2 นิ้ว]' placeholder text from the photo frame
    tbl_bio = doc.tables[-1]
    cell_photo = tbl_bio.cell(0, 1)
    for nt in cell_photo.tables:
        for row in nt.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    for r in p.runs:
                        if "รูปถ่าย" in r.text or "[" in r.text:
                            r.text = ""
                    if "รูปถ่าย" in p.text or "[" in p.text:
                        p.text = ""
    print("Cleared photo placeholder text from Author Bio.")

    # b. Remove trailing empty paragraph (w:spacing w:before="4800") that causes empty page 129
    p_last = doc.paragraphs[-1]
    if not p_last.text.strip():
        p_last._p.getparent().remove(p_last._p)
        print("Removed trailing empty paragraph at the end of book.")

    doc.save(NEW_V5_DOCX)
    print(f"Saved updated docx: {NEW_V5_DOCX}")

    print("Exporting to PDF via Word COM...")
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        wdoc = word.Documents.Open(os.path.abspath(NEW_V5_DOCX))
        # DO NOT update TOC fields, so page numbers in TOC stay unchanged as requested!
        wdoc.Repaginate()
        wdoc.Save()
        wdoc.ExportAsFixedFormat(os.path.abspath(NEW_V5_PDF), 17) # 17 = wdExportFormatPDF
        wdoc.Close(False)
        print(f"Exported PDF successfully: {NEW_V5_PDF}")
    finally:
        word.Quit()

if __name__ == "__main__":
    update_book()
