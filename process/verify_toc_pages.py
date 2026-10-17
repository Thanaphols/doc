import sys
import docx
import pymupdf
import re

sys.stdout.reconfigure(encoding='utf-8')

pdf = pymupdf.open('doc/output/Chatshop_ปริญญานิพนธ์ฉบับสมบูรณ์_v7.pdf')
doc = docx.Document('doc/output/Chatshop_ปริญญานิพนธ์ฉบับสมบูรณ์_v7.docx')

print(f"Total PDF pages: {len(pdf)}")

# Body starts at PDF page index 14 (1-indexed page 15 = Body Page 1)
def pdf_idx_to_body_page(idx):
    return idx - 14 + 1

# Prepare text of all body pages (from index 14 to end)
body_pages_text = []
for idx in range(14, len(pdf)):
    page_text = pdf[idx].get_text()
    body_pages_text.append((idx, pdf_idx_to_body_page(idx), page_text))

def parse_items(start_p, end_p, section_type):
    items = []
    for i in range(start_p, end_p + 1):
        if i >= len(doc.paragraphs): break
        t = doc.paragraphs[i].text.strip()
        if not t: continue
        # Lines might contain multiple entries separated by newline
        lines = t.split('\n')
        for line in lines:
            line = line.strip()
            if not line: continue
            if line.startswith("สารบัญ") or line.startswith("หน้า") or line.startswith("ตาราง\tหน้า") or line.startswith("รูป\tหน้า"):
                continue
            parts = line.rsplit('\t', 1)
            if len(parts) == 2:
                title = parts[0].strip()
                page_str = parts[1].strip()
                items.append((section_type, title, page_str))
    return items

toc_main = parse_items(87, 172, "สารบัญ")
toc_tables = parse_items(173, 201, "สารบัญตาราง")
toc_figures = parse_items(202, 276, "สารบัญรูป")

all_items = toc_main + toc_tables + toc_figures
print(f"Total entries: สารบัญ={len(toc_main)}, ตาราง={len(toc_tables)}, รูป={len(toc_figures)}, รวม={len(all_items)}\n")

def find_actual_page(item_type, title):
    if item_type == "สารบัญตาราง":
        m = re.match(r'(ตารางที่\s*\d+\.\d+)', title)
        if m:
            prefix = m.group(1).replace(' ', '')
            for idx, b_page, text in body_pages_text:
                compact_text = text.replace(' ', '')
                if prefix in compact_text:
                    return b_page
    elif item_type == "สารบัญรูป":
        m = re.match(r'(รูปที่\s*[\dก]\.\d+)', title)
        if m:
            prefix = m.group(1).replace(' ', '')
            for idx, b_page, text in body_pages_text:
                compact_text = text.replace(' ', '')
                if prefix in compact_text:
                    return b_page
    else: # สารบัญ
        if title.startswith("บทที่"):
            chap_num = title.split()[1]
            chap_str = f"บทที่ {chap_num}".replace(' ', '')
            for idx, b_page, text in body_pages_text:
                lines = [l.strip().replace(' ', '') for l in text.split('\n') if l.strip()]
                if len(lines) > 0 and lines[0] == chap_str:
                    return b_page
                elif len(lines) > 1 and lines[1] == chap_str:
                    return b_page
        elif "บรรณานุกรม" in title:
            for idx, b_page, text in body_pages_text:
                lines = [l.strip() for l in text.split('\n') if l.strip()]
                if "บรรณานุกรม" in lines[:3]:
                    return b_page
        elif "ภาคผนวก ก" in title:
            for idx, b_page, text in body_pages_text:
                if "ก.1 ข้อกำหนดและผู้ใช้งานแอปพลิเคชัน" in text:
                    return b_page
        elif "ประวัติผู้เขียน" in title:
            for idx, b_page, text in body_pages_text:
                if "ประวัติผู้เขียน" in text:
                    return b_page
        else: # Section / subsection
            m = re.match(r'^([\dก-ฮ]+\.\d+(\.\d+)?)\s*(.*)', title)
            if m:
                prefix = m.group(1)
                rest = m.group(3).strip()
                # Clean rest for searching
                clean_rest = re.sub(r'[\s\(\)\[\]]+', '', rest)[:15]
                for idx, b_page, text in body_pages_text:
                    compact_text = re.sub(r'[\s\(\)\[\]]+', '', text)
                    if (prefix + clean_rest) in compact_text or (prefix in text and clean_rest in compact_text):
                        return b_page
            else:
                clean_title = re.sub(r'[\s\(\)\[\]]+', '', title)[:20]
                for idx, b_page, text in body_pages_text:
                    compact_text = re.sub(r'[\s\(\)\[\]]+', '', text)
                    if clean_title in compact_text:
                        return b_page
    return None

matches = []
diffs = []

for item_type, title, stated in all_items:
    # If stated is roman/thai page number (like ก, ข, ค...), check separately
    if stated in ['ก','ข','ค','ง','จ','ฉ','ช','ซ','ฌ','ญ','ฎ','ฏ']:
        matches.append((item_type, title, stated, stated))
        continue
        
    actual = find_actual_page(item_type, title)
    if actual is not None:
        if str(actual) == str(stated):
            matches.append((item_type, title, stated, actual))
        else:
            diffs.append((item_type, title, stated, actual))
    else:
        diffs.append((item_type, title, stated, "NOT_FOUND"))

print(f"=== SUMMARY ===")
print(f"Exact Matches: {len(matches)} / {len(all_items)}")
print(f"Differences  : {len(diffs)} / {len(all_items)}")

if diffs:
    print(f"\n{'SECTION':<14} | {'STATED':<6} | {'ACTUAL':<6} | {'STATUS':<10} | {'TITLE'}")
    print("-" * 85)
    for t_type, title, stated, actual in diffs:
        status = "DIFF" if actual != "NOT_FOUND" else "NOT_FOUND"
        print(f"{t_type:<14} | {stated:<6} | {str(actual):<6} | {status:<10} | {title}")
