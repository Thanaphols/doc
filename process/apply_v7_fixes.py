import os
import sys
from PIL import Image, ImageOps, ImageEnhance
import docx
from docx.shared import Pt, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsdecls
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.stdout.reconfigure(encoding='utf-8')

DOCX_IN = r"D:\Thesis\doc\output\Chatshop_ปริญญานิพนธ์ฉบับสมบูรณ์_v7_user_saved.docx"
DOCX_OUT = r"D:\Thesis\doc\output\Chatshop_ปริญญานิพนธ์ฉบับสมบูรณ์_v7.docx"
TEMP_DIR = r"D:\Thesis\doc\process\temp\bw_ch2"
os.makedirs(TEMP_DIR, exist_ok=True)

print("Loading document:", DOCX_IN)
doc = docx.Document(DOCX_IN)

# -------------------------------------------------------------
# 1. PROCESS AND UPDATE CH2 IMAGES
# -------------------------------------------------------------
print("\n--- 1. Processing and updating Ch2 images ---")

def process_ch2_img(in_path, out_path):
    im = Image.open(in_path).convert('RGBA')
    bg = Image.new('RGBA', im.size, (255, 255, 255, 255))
    comp = Image.alpha_composite(bg, im).convert('RGB')
    
    gray = ImageOps.grayscale(comp)
    mask = gray.point(lambda p: 255 if p < 248 else 0)
    bbox = mask.getbbox()
    if bbox:
        pad = 25
        w, h = comp.size
        bbox = (max(0, bbox[0]-pad), max(0, bbox[1]-pad), min(w, bbox[2]+pad), min(h, bbox[3]+pad))
        comp = comp.crop(bbox)
        gray = gray.crop(bbox)
    
    def tone_map(p):
        if p > 240:
            return 255
        elif p < 130:
            return int(p * 0.4) # deepen dark text and lines to black
        elif p < 190:
            return int(p * 0.75)
        else:
            return int(200 + (p - 190) * 0.9) # light tints stay light gray
            
    bw = gray.point(tone_map)
    bw = ImageEnhance.Sharpness(bw).enhance(2.0)
    bw.save(out_path, optimize=True)
    return bw.size

# Target width in EMUs: 13.5 cm = 4,860,000 EMUs (fits comfortably within 14.65 cm body width)
TARGET_CX = 4860000

for i in range(1, 20):
    src_img = f"D:/Thesis/doc/input/figures/ch2/2.{i}.png"
    out_img = f"{TEMP_DIR}/2.{i}.png"
    if not os.path.exists(src_img):
        print(f"Warning: {src_img} not found")
        continue
    
    img_w, img_h = process_ch2_img(src_img, out_img)
    target_cy = int(TARGET_CX * img_h / img_w)
    
    # Table index for Figure 2.i is table i + 2 (since Table 3 is Fig 2.1, Table 4 is Fig 2.2, etc.)
    tbl_idx = i + 2
    table = doc.tables[tbl_idx]
    cell = table.cell(0, 0)
    
    # Update blip embed binary
    blips = cell._element.xpath('.//a:blip')
    if blips:
        rId = blips[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
        part = doc.part.related_parts[rId]
        with open(out_img, 'rb') as f:
            part._blob = f.read()
            
    # Update extents
    for ext in cell._element.xpath('.//wp:extent'):
        ext.set('cx', str(TARGET_CX))
        ext.set('cy', str(target_cy))
    for ext in cell._element.xpath('.//a:ext'):
        ext.set('cx', str(TARGET_CX))
        ext.set('cy', str(target_cy))
        
    print(f"Updated Figure 2.{i} in Table {tbl_idx}: {img_w}x{img_h} -> cx={TARGET_CX}, cy={target_cy}")

# -------------------------------------------------------------
# 2. FIX PAGE 90: REMOVE EMPTY LINES & PAGEBREAKBEFORE ON 4.18
# -------------------------------------------------------------
print("\n--- 2. Fixing Page 90 ---")
# Find paragraph 4.18 (the actual section heading in Chapter 4 body, not the chapter overview)
p_418 = None
p_418_idx = None
for idx, p in enumerate(doc.paragraphs):
    if "การเปรียบเทียบการทำงานก่อนและหลังใช้แอปพลิเคชัน Chatshop" in p.text and idx > 950:
        p_418 = p
        p_418_idx = idx
        break

if p_418 is not None:
    print(f"Found 4.18 at index {p_418_idx}: {p_418.text[:60]}")
    # Remove pageBreakBefore from 4.18
    pPr = p_418._element.pPr
    if pPr is not None:
        pbb = pPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pageBreakBefore')
        if pbb is not None:
            pPr.remove(pbb)
            print("Removed pageBreakBefore from 4.18")
            
    # Look backwards from p_418_idx and remove empty paragraphs before it
    # P[982] to P[989] were empty paragraphs
    curr_idx = p_418_idx - 1
    removed_count = 0
    while curr_idx >= 0 and doc.paragraphs[curr_idx].text.strip() == "":
        # Check that this empty paragraph doesn't contain sectPr
        p_empty = doc.paragraphs[curr_idx]
        if "sectPr" not in p_empty._element.xml:
            p_elem = p_empty._element
            p_elem.getparent().remove(p_elem)
            removed_count += 1
            curr_idx -= 1
        else:
            break
    print(f"Removed {removed_count} empty paragraphs before 4.18")
else:
    print("Warning: Could not find 4.18 heading!")

# -------------------------------------------------------------
# 3. FIX APPENDIX A: INSERT DIVIDER PAGE FOR ภาคผนวก ก
# -------------------------------------------------------------
print("\n--- 3. Fixing Appendix A divider page ---")

# Find the start of Section 15 / manual content
p_manual_start = None
p_manual_idx = None
for idx, p in enumerate(doc.paragraphs):
    if "ภาคผนวก ก" in p.text and "คู่มือการใช้งาน" in p.text and idx > 1000:
        p_manual_start = p
        p_manual_idx = idx
        break

if p_manual_start is not None:
    print(f"Found manual start at index {p_manual_idx}: {repr(p_manual_start.text[:50])}")
    
    # We want:
    # 1. Divider Page for ภาคผนวก ก:
    #    - Blank lines to push text to middle of page
    #    - "ภาคผนวก ก" (center, bold, 18pt)
    #    - "คู่มือการใช้งานแอปพลิเคชัน Chatshop" (center, bold, 18pt)
    # 2. Followed by a page break so p_manual_start begins on the next page!
    
    # Ensure p_manual_start has pageBreakBefore so it starts on next page after divider
    pPr = p_manual_start._element.get_or_add_pPr()
    pbb = pPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pageBreakBefore')
    if pbb is None:
        new_pbb = OxmlElement('w:pageBreakBefore')
        pPr.append(new_pbb)
        print("Added pageBreakBefore to manual content start")
        
    # Create divider page paragraphs before p_manual_start
    # In Section 15, page 1 is different_first_page so it will have NO header/page number!
    parent = p_manual_start._element.getparent()
    
    def create_p(text="", bold=False, size_pt=16, align=WD_ALIGN_PARAGRAPH.CENTER):
        p = OxmlElement('w:p')
        pPr = OxmlElement('w:pPr')
        jc = OxmlElement('w:jc')
        jc.set(qn('w:val'), 'center' if align == WD_ALIGN_PARAGRAPH.CENTER else 'left')
        pPr.append(jc)
        p.append(pPr)
        
        if text:
            r = OxmlElement('w:r')
            rPr = OxmlElement('w:rPr')
            rFonts = OxmlElement('w:rFonts')
            rFonts.set(qn('w:ascii'), 'TH SarabunPSK')
            rFonts.set(qn('w:hAnsi'), 'TH SarabunPSK')
            rFonts.set(qn('w:cs'), 'TH SarabunPSK')
            rPr.append(rFonts)
            
            if bold:
                b = OxmlElement('w:b')
                bCs = OxmlElement('w:bCs')
                rPr.append(b)
                rPr.append(bCs)
                
            sz = OxmlElement('w:sz')
            sz.set(qn('w:val'), str(int(size_pt * 2)))
            szCs = OxmlElement('w:szCs')
            szCs.set(qn('w:val'), str(int(size_pt * 2)))
            rPr.append(sz)
            rPr.append(szCs)
            
            r.append(rPr)
            t = OxmlElement('w:t')
            t.text = text
            r.append(t)
            p.append(r)
        return p

    # Insert 9 empty lines
    for _ in range(9):
        p_blank = create_p("", size_pt=16)
        parent.insert(parent.index(p_manual_start._element), p_blank)
        
    # Insert "ภาคผนวก ก"
    p_title1 = create_p("ภาคผนวก ก", bold=True, size_pt=18)
    parent.insert(parent.index(p_manual_start._element), p_title1)
    
    # Insert "คู่มือการใช้งานแอปพลิเคชัน Chatshop"
    p_title2 = create_p("คู่มือการใช้งานแอปพลิเคชัน Chatshop", bold=True, size_pt=18)
    parent.insert(parent.index(p_manual_start._element), p_title2)
    
    print("Inserted Appendix A divider page paragraphs successfully")
else:
    print("Warning: Could not find manual start paragraph!")

# -------------------------------------------------------------
# SAVE DOCUMENT
# -------------------------------------------------------------
print(f"\nSaving updated document to: {DOCX_OUT}")
doc.save(DOCX_OUT)
print("Saved successfully!")
