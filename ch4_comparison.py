"""Section 4.18 (before/after comparison) and the matching lines in chapter 5.
Process-based comparison; every score is from the real 15-person survey in 4.17.
No measured time/volume figures exist, so none are given (rule: no invented numbers)."""
import copy

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm

HEADING = "การเปรียบเทียบการทำงานก่อนและหลังใช้แอปพลิเคชัน Chatshop"
INTRO = ("ผู้จัดทำเปรียบเทียบขั้นตอนการทำงานของเจ้าหน้าที่ร้านค้าออนไลน์แบบเดิมที่ใช้แอปพลิเคชันของแต่ละแพลตฟอร์มแยกกัน "
         "กับการทำงานผ่านแอปพลิเคชัน Chatshop โดยพิจารณาจากความสามารถของระบบที่ผ่านการทดสอบในหัวข้อ 4.1 – 4.16 "
         "ร่วมกับผลการประเมินความพึงพอใจในหัวข้อ 4.17 ดังตารางที่ 4.23")
CAPTION = "การเปรียบเทียบการทำงานก่อนและหลังใช้แอปพลิเคชัน Chatshop"
ROWS = [
    ("ด้าน", "การทำงานแบบเดิม", "เมื่อใช้แอปพลิเคชัน Chatshop"),
    ("ช่องทางรับข้อความ", "เปิดแอปพลิเคชัน LINE Official Account, Facebook Messenger, Instagram และหน้าเว็บไซต์แยกกัน",
     "ข้อความจากทั้ง 4 ช่องทางรวมอยู่ในกล่องข้อความเดียว"),
    ("การรับรู้ข้อความใหม่", "ต้องเปิดตรวจทีละแอปพลิเคชัน",
     "ได้รับ Push Notification และกล่องข้อความอัปเดตแบบเรียลไทม์"),
    ("แชทที่ค้างตอบ", "ไม่มีการเตือน ลูกค้าอาจรอนานโดยไม่มีผู้ใดสังเกต",
     "ระบบแจ้งเตือนเมื่อแชทค้างตอบเกินเวลาที่ร้านกำหนด"),
    ("การตอบซ้ำซ้อน", "เจ้าหน้าที่หลายคนอาจตอบลูกค้าคนเดียวกันพร้อมกัน",
     "ล็อกบทสนทนาให้เจ้าหน้าที่คนแรกที่ตอบ และมอบหมายผู้รับผิดชอบได้"),
    ("คำถามที่พบบ่อย", "เจ้าหน้าที่ต้องพิมพ์ตอบเองทุกข้อความ",
     "แชทบอทแนะนำคำตอบหรือตอบลูกค้าเองตามโหมดที่ร้านเลือก และส่งต่อเจ้าหน้าที่เมื่อตอบไม่ได้"),
    ("การบันทึกคำสั่งซื้อ", "คัดลอกข้อมูลจากแชทไปจดในสมุดหรือโปรแกรมอื่น",
     "เปิดคำสั่งซื้อจากหน้าต่างสนทนา ระบบคำนวณยอดรวมและสร้างเลขที่คำสั่งซื้อให้อัตโนมัติ"),
    ("ข้อมูลลูกค้า", "กระจายอยู่ในแชทของแต่ละแพลตฟอร์ม",
     "แผ่นข้อมูลลูกค้าแสดงที่อยู่ บันทึก แท็ก และประวัติคำสั่งซื้อในหน้าต่างสนทนา"),
]
WIDTHS_CM = (3.3, 5.6, 5.75)
DISCUSSION = [
    "จากตารางที่ 4.23 ประโยชน์ของแอปพลิเคชัน Chatshop เมื่อเทียบกับการทำงานแบบเดิม สรุปได้ 3 ด้าน ดังนี้",
    "4.18.1 ตอบแชทได้มากขึ้น เมื่อข้อความจากทุกช่องทางรวมอยู่ที่เดียวและแชทบอทช่วยตอบคำถามที่พบบ่อย "
    "เจ้าหน้าที่ใช้เวลากับการสลับแอปพลิเคชันน้อยลงและรับมือแชทได้มากขึ้นในเวลาเท่าเดิม สอดคล้องกับผลการประเมิน"
    "ข้อความสะดวกในการรวมข้อความจากทุกช่องทางไว้ในกล่องข้อความเดียว ที่ได้ค่าเฉลี่ยสูงสุดเท่ากับ 4.27",
    "4.18.2 ลดความผิดพลาด การล็อกบทสนทนาช่วยป้องกันการตอบลูกค้าซ้ำซ้อน และการเปิดคำสั่งซื้อจากแชทที่ระบบ"
    "คำนวณยอดรวมให้เองช่วยลดการคัดลอกข้อมูลผิด ผู้ประเมินให้คะแนนข้อระบบช่วยลดข้อผิดพลาดในการตอบแชทและ"
    "การบันทึกคำสั่งซื้อเฉลี่ย 4.13 และข้อความสะดวกและความถูกต้องในการเปิดและจัดการคำสั่งซื้อเฉลี่ย 4.07",
    "4.18.3 ประหยัดเวลา เจ้าหน้าที่ไม่ต้องเปิดหลายแอปพลิเคชันเพื่อตรวจข้อความ และได้รับการแจ้งเตือนทันทีเมื่อมี"
    "ข้อความใหม่หรือแชทค้างตอบ ผู้ประเมินให้คะแนนข้อระบบช่วยลดภาระและลดเวลาในการสลับหลายแอปพลิเคชัน"
    "เฉลี่ย 4.20 ซึ่งอยู่ในระดับดี",
    "อย่างไรก็ตาม การเปรียบเทียบนี้อิงจากขั้นตอนการทำงานและผลการประเมินความพึงพอใจ ยังไม่ได้วัดเวลาในการตอบ"
    "หรือจำนวนแชทก่อนและหลังใช้ระบบเป็นตัวเลข ซึ่งควรเก็บข้อมูลจากหน้ารายงานของระบบในการใช้งานจริงต่อไป",
]
CH5_SUMMARY = ("5.1.7 เมื่อเปรียบเทียบกับการทำงานแบบเดิม แอปพลิเคชันช่วยให้เจ้าหน้าที่รับมือแชทได้มากขึ้น "
               "จากการรวมข้อความทุกช่องทางไว้ที่เดียวและให้แชทบอทช่วยตอบ ลดความผิดพลาดจากการตอบซ้ำซ้อนและ"
               "การคัดลอกข้อมูลคำสั่งซื้อ และลดเวลาในการสลับแอปพลิเคชัน ตามหัวข้อ 4.18")
CH5_MEASURE = ("5.4.7 เก็บข้อมูลเวลาในการตอบครั้งแรกและจำนวนแชทที่ตอบได้จากหน้ารายงานของระบบ ทั้งก่อนและหลัง"
               "ใช้งาน เพื่อวัดผลการเปรียบเทียบเป็นตัวเลข")


def clone_after(anchor_el, template_p, text):
    """New paragraph formatted like template_p, placed after anchor_el."""
    el = copy.deepcopy(template_p._p)
    anchor_el.addnext(el)
    p = docx.text.paragraph.Paragraph(el, template_p._parent)
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)
    for fld in el.findall(qn("w:fldSimple")):
        el.remove(fld)
    if el.pPr is not None:
        for s in el.pPr.findall(qn("w:sectPr")):
            el.pPr.remove(s)
    p.runs[0].text = text
    return p


def _set_text(p, text):
    runs = [r for r in p.runs if r.text]
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def blank_before_first_table(doc):
    """Table 4.1 sat right under the chapter's topic list: one blank line above its caption."""
    cap = next(p for p in doc.paragraphs if p.text.startswith("ตารางที่"))
    prev = cap._p.getprevious()
    if prev is not None and docx.text.paragraph.Paragraph(prev, doc).text.strip():
        cap._p.addprevious(OxmlElement("w:p"))


def add_comparison(doc):
    """Append section 4.18 + table 4.23 at the end of chapter 4 (standalone chapter doc)."""
    ps = doc.paragraphs
    eval_h2 = next(p for p in ps if p.style.name == "Heading 2" and p.text.strip().startswith("แบบประเมิน"))
    body_p = next(p for p in ps if p.text.startswith("\tผู้จัดทำได้นำแอปพลิเคชัน Chatshop"))
    item_p = next(p for p in ps if p.text.strip().startswith("4.17.1 "))
    topic = next(p for p in ps if p.text.strip() == "4.17 แบบประเมินความพึงพอใจในการใช้งานแอปพลิเคชัน")
    clone_after(topic._p, topic, "4.18 " + HEADING)

    anchor = [el for el in doc.element.body if el.tag == qn("w:p")][-1]
    h = clone_after(anchor, eval_h2, HEADING)
    h.paragraph_format.page_break_before = True
    anchor = clone_after(h._p, body_p, "\t" + INTRO)._p

    cap_src = next(p for p in doc.paragraphs if p.text.startswith("ตารางที่ 4.") and "สรุปผลการประเมิน" in p.text)
    cap = copy.deepcopy(cap_src._p)
    anchor.addnext(cap)
    for t in cap.iter(qn("w:t")):
        if t.text == "22":
            t.text = "23"  # cached SEQ result
        elif "สรุปผลการประเมิน" in (t.text or ""):
            t.text = "  " + CAPTION
    tbl = copy.deepcopy(cap_src._p.getnext())  # table 4.22: same borders, fonts, header row
    cap.addnext(tbl)
    t = docx.table.Table(tbl, doc)
    while len(t.rows) < len(ROWS):
        tbl.append(copy.deepcopy(t.rows[1]._tr))
    while len(t.rows) > len(ROWS):
        tbl.remove(t.rows[-1]._tr)
    for gc in tbl.tblGrid.findall(qn("w:gridCol"))[3:]:
        tbl.tblGrid.remove(gc)
    for row in t.rows:
        for tc in row._tr.tc_lst[3:]:
            row._tr.remove(tc)
    for gc, w in zip(tbl.tblGrid.findall(qn("w:gridCol")), WIDTHS_CM):
        gc.set(qn("w:w"), str(round(w * 567)))
    t.autofit = False
    for r_i, (row, vals) in enumerate(zip(t.rows, ROWS)):
        for c_i, (cell, val) in enumerate(zip(row.cells, vals)):
            cell.width = Cm(WIDTHS_CM[c_i])
            for extra in cell.paragraphs[1:]:
                extra._p.getparent().remove(extra._p)
            _set_text(cell.paragraphs[0], val)
            if r_i and c_i:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    anchor = clone_after(tbl, body_p, "")._p  # blank line under the table
    for text in DISCUSSION:
        anchor = clone_after(anchor, item_p if text.startswith("4.18.") else body_p, "\t" + text)._p


def update_ch5(doc):
    """5.1.7 summary after 5.1.6; drop 5.4.6, renumber 5.4.7, then 5.4.7 = measure the gain."""
    def para(prefix):
        return next(p for p in doc.paragraphs if p.text.strip().startswith(prefix))
    p516 = para("5.1.6 ")
    clone_after(p516._p, p516, "\t" + CH5_SUMMARY)
    para("5.4.6 ")._p.getparent().remove(para("5.4.6 ")._p)
    p547 = para("5.4.7 ")
    r = next(r for r in p547.runs if "5.4.7" in r.text)
    r.text = r.text.replace("5.4.7", "5.4.6", 1)
    clone_after(p547._p, p547, "\t" + CH5_MEASURE)
