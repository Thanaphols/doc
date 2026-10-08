import os, sys, shutil, docx
from docx.oxml.ns import qn

sys.path.insert(0, '.agent/skills/rmuti-thesis/scripts')
import thesis_build

ch3_path = os.path.join('output', 'Chatshop_บทที่3_วิธีดำเนินงาน.docx')
backup_path = os.path.join('output', 'Chatshop_บทที่3_วิธีดำเนินงาน_backup.docx')

if not os.path.exists(backup_path):
    shutil.copy2(ch3_path, backup_path)
    print(f'Backed up {ch3_path} to {backup_path}')

doc = docx.Document(ch3_path)
fixed_runs = 0

for p in doc.paragraphs:
    for r in p.runs:
        rPr = r._r.rPr
        if rPr is None or rPr.find(qn('w:cs')) is None:
            thesis_build.mark_thai(r)
            fixed_runs += 1

for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    rPr = r._r.rPr
                    if rPr is None or rPr.find(qn('w:cs')) is None:
                        thesis_build.mark_thai(r)
                        fixed_runs += 1

doc.save(ch3_path)
print(f'Successfully fixed {fixed_runs} runs with mark_thai in {ch3_path}')
